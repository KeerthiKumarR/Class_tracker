"""Extracts text from PDF, DOCX, XLSX, or text files and converts syllabus to structured curriculum JSON via Groq LLM."""

import os
import sys
import json
import re
from typing import List, Dict, Any, Optional
from groq import Groq

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(HERE)


def extract_text_from_file(file_path: str) -> str:
    """Extracts raw text content from PDF, DOCX, XLSX, or plain text files."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            pages_text = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    pages_text.append(f"--- Page {i+1} ---\n{text}")
            return "\n\n".join(pages_text)
        except Exception as e:
            raise RuntimeError(f"Failed to read PDF: {e}")

    elif ext in (".docx", ".doc"):
        try:
            import docx
            doc = docx.Document(file_path)
            lines = []
            for p in doc.paragraphs:
                if p.text.strip():
                    lines.append(p.text)
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)
        except Exception as e:
            raise RuntimeError(f"Failed to read DOCX: {e}")

    elif ext in (".xlsx", ".xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            lines = []
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                lines.append(f"=== Sheet: {sheetname} ===")
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if row_vals:
                        lines.append(" | ".join(row_vals))
            return "\n".join(lines)
        except Exception as e:
            raise RuntimeError(f"Failed to read Excel file: {e}")

    else:
        # Assume plain text or markdown
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception as e:
            raise RuntimeError(f"Failed to read text file: {e}")


def try_parse_structured_table(raw_text: str, course_code: str = "CRS") -> List[Dict[str, Any]]:
    """Fast, deterministic parser for structured syllabi with lecture numbers (e.g. LEC-1, Lecture 1).
    Consumes 0 Groq tokens and is immune to rate limits.
    """
    clean_code = (course_code or "CRS").upper().strip()
    
    # Pattern for LEC-X
    lec_pattern = re.compile(
        r'(?:^|\n)\s*(?:(\d+)\s+)?LEC-(\d+)\s+([^\n•]+?)(?=\s*•|\s*\n|$)(.*?)(?=(?:\n\s*(?:\d+\s+)?LEC-\d+)|\Z)',
        re.DOTALL
    )
    
    matches = list(lec_pattern.finditer(raw_text))
    if len(matches) < 3:
        # Also try "Lecture X" or "Unit X"
        alt_pattern = re.compile(
            r'(?:^|\n)\s*(?:(\d+)\s+)?(?:Lecture|Unit|Module)\s*(\d+)[:\s\-]+([^\n•\|]+?)(?=\s*[•\|]|\s*\n|$)(.*?)(?=(?:\n\s*(?:\d+\s+)?(?:Lecture|Unit|Module)\s*\d+)|\Z)',
            re.DOTALL | re.IGNORECASE
        )
        matches = list(alt_pattern.finditer(raw_text))
        if len(matches) < 3:
            return []
    
    units = []
    for match in matches:
        week_str, lec_num, topic, body = match.groups()
        try:
            week = int(week_str) if week_str and week_str.isdigit() else ((int(lec_num) - 1) // 5 + 1)
        except Exception:
            week = 1
        topic = topic.strip()
        
        # Bullets for content
        bullets = [b.strip() for b in re.findall(r'[•\-\*]\s*([^\n•\-\*]+)', body)]
        
        # Numbered items for classwork / homework
        numbered = [n.strip() for n in re.findall(r'(?:^|\n)\s*\d+[\.\)]\s*([^\n]+(?:\n(?!\s*\d+[\.\)]|\s*[•\-\*])[^\n]+)*)', body)]
        
        content = bullets
        classwork = []
        homework = []
        
        if numbered:
            mid = len(numbered) // 2
            classwork = numbered[:mid] if mid > 0 else numbered
            homework = numbered[mid:] if mid > 0 else []
        elif not bullets:
            lines = [l.strip() for l in body.split('\n') if l.strip()]
            content = lines[:4]
            classwork = lines[4:7]
            homework = lines[7:10]
            
        units.append({
            "unit_id": f"{clean_code}-{lec_num}",
            "week": week,
            "topic": topic,
            "content": content,
            "classwork": classwork,
            "homework": homework,
        })
        
    return units


def parse_syllabus_to_curriculum(
    raw_text: str,
    course_name: str = "Course",
    course_code: str = "CRS",
    groq_api_key: Optional[str] = None,
    groq_model: str = "openai/gpt-oss-120b",
) -> List[Dict[str, Any]]:
    """Converts syllabus text into standard curriculum JSON units.
    Tries fast deterministic structured table extraction first, then falls back to Groq LLM.
    """
    if not raw_text or not raw_text.strip():
        raise ValueError("Syllabus text is empty.")

    # 1. Try fast structured parsing first (0 Groq tokens, instant, no rate limits)
    structured = try_parse_structured_table(raw_text, course_code)
    if len(structured) >= 3:
        print(f"[PARSER] Deterministically extracted {len(structured)} units without LLM.")
        return structured

    if not groq_api_key:
        groq_api_key = os.environ.get("GROQ_API_KEY")

    if not groq_api_key:
        raise ValueError("GROQ_API_KEY is not configured.")

    client = Groq(api_key=groq_api_key, timeout=45.0)

    # Truncate text to stay well within Groq TPM limits
    cleaned_text = raw_text[:8000]

    system_prompt = f"""You are a specialized Curriculum Extraction AI.
Extract syllabus text for course "{course_name}" ({course_code}) into a valid JSON array of lecture units.
Every unit represents a 1-class session ({course_code}-1, {course_code}-2, ...).
Schema:
[{{
  "unit_id": "{course_code}-1",
  "week": 1,
  "topic": "Topic Name",
  "content": ["Concept 1"],
  "classwork": ["Problem 1"],
  "homework": ["Problem 1"]
}}]
Return strictly valid JSON array [ ... ]."""

    user_prompt = f"Extract all lecture units from this syllabus text:\n\n{cleaned_text}"

    # Try models in order: 120b -> 20b -> qwen
    models_to_try = [groq_model, "openai/gpt-oss-120b", "openai/gpt-oss-20b", "qwen/qwen3.8-27b"]
    seen_models = set()
    response = None

    for m in models_to_try:
        if not m or m in seen_models:
            continue
        seen_models.add(m)
        tok_limit = 1000 if "qwen" in m else 2500
        try:
            response = client.chat.completions.create(
                model=m,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
                temperature=0.1,
                max_tokens=tok_limit,
            )
            break
        except Exception as ge:
            print(f"[WARN] Model {m} failed: {ge}, trying fallback...")
            continue

    if not response:
        raise RuntimeError("All Groq models encountered rate limits or errors.")

    content = response.choices[0].message.content or ""

    # Parse JSON from response
    clean_json_str = content.strip()
    if "```json" in clean_json_str:
        clean_json_str = clean_json_str.split("```json")[1].split("```")[0].strip()
    elif "```" in clean_json_str:
        clean_json_str = clean_json_str.split("```")[1].split("```")[0].strip()

    try:
        parsed = json.loads(clean_json_str)
    except Exception as e:
        # Fallback regex extraction of JSON array or object
        match = re.search(r"(\[.*\]|\{.*\})", clean_json_str, re.DOTALL)
        if match:
            parsed = json.loads(match.group(1))
        else:
            raise ValueError(f"Could not parse LLM output as JSON: {e}\nRaw output: {content[:300]}")

    if isinstance(parsed, dict):
        for k in ("units", "curriculum", "topics", "data"):
            if k in parsed and isinstance(parsed[k], list):
                parsed = parsed[k]
                break
        else:
            parsed = [parsed]

    # Normalize units
    normalized_units = []
    for idx, u in enumerate(parsed, start=1):
        uid = u.get("unit_id") or f"{course_code}-{idx}"
        topic = u.get("topic") or f"Module {idx}"
        week = u.get("week") or ((idx - 1) // 3 + 1)
        content_items = u.get("content") or []
        if isinstance(content_items, str):
            content_items = [c.strip() for c in content_items.split("\n") if c.strip()]
        classwork = u.get("classwork") or []
        if isinstance(classwork, str):
            classwork = [c.strip() for c in classwork.split("\n") if c.strip()]
        homework = u.get("homework") or []
        if isinstance(homework, str):
            homework = [h.strip() for h in homework.split("\n") if h.strip()]

        normalized_units.append({
            "unit_id": str(uid).strip(),
            "week": int(week) if str(week).isdigit() else 1,
            "topic": str(topic).strip(),
            "content": content_items,
            "classwork": classwork,
            "homework": homework,
        })

    return normalized_units


def parse_multi_subject_syllabus(
    raw_text: str,
    available_tabs: Optional[List[str]] = None,
    groq_api_key: Optional[str] = None,
    groq_model: str = "openai/gpt-oss-120b",
) -> List[Dict[str, Any]]:
    """Analyzes a multi-subject document, detects each distinct course (e.g. DAA, OS, DSA),
    maps to available spreadsheet tabs, and extracts structured curriculum units for each course.
    """
    if not raw_text or not raw_text.strip():
        raise ValueError("Syllabus text is empty.")

    # 1. Fast deterministic structured table parsing (0 tokens, instant, no rate limits)
    text_lower = raw_text.lower()
    code_cand = "CRS"
    name_cand = "Course"
    if "dsa" in text_lower or "data structure" in text_lower:
        code_cand = "DSA"
        name_cand = "Data Structures & Algorithms"
    elif "daa" in text_lower or "design & analysis" in text_lower:
        code_cand = "DAA"
        name_cand = "Design & Analysis of Algorithm (DAA)"
    elif "os" in text_lower or "operating system" in text_lower:
        code_cand = "OS"
        name_cand = "Operating Systems"

    matched_tab = code_cand
    if available_tabs:
        for t in available_tabs:
            if t.lower() in text_lower or code_cand.lower() in t.lower():
                matched_tab = t
                break

    structured = try_parse_structured_table(raw_text, code_cand)
    if len(structured) >= 3:
        print(f"[PARSER] Deterministically extracted {len(structured)} structured units for {code_cand}.")
        return [{
            "course_name": name_cand,
            "course_code": code_cand,
            "matched_tab": matched_tab,
            "units": structured
        }]

    if not groq_api_key:
        groq_api_key = os.environ.get("GROQ_API_KEY")

    if not groq_api_key:
        raise ValueError("GROQ_API_KEY is not configured.")

    client = Groq(api_key=groq_api_key, timeout=60.0)
    cleaned_text = raw_text[:8000]

    tab_hint = f"\nAvailable Google Spreadsheet Tabs: {available_tabs}" if available_tabs else ""

    system_prompt = f"""You are an Expert Multi-Course Curriculum Ingestion AI.
The user provided syllabus text that may contain ONE, TWO, THREE, OR MORE distinct subjects/courses
(for example: DAA, Operating Systems, Data Structures, etc.).{tab_hint}

Your job:
1. Detect each distinct subject/course in the document.
2. For each course, extract its structured syllabus units at the INDIVIDUAL 1-CLASS/1-LECTURE LEVEL (e.g. 1 hour class session).
   - DO NOT lump an entire week of multiple subtopics and problems into one giant unit!
   - Split each week into 3 to 5 separate lecture units (e.g. {available_tabs[0] if available_tabs else 'CRS'}-1, {available_tabs[0] if available_tabs else 'CRS'}-2, etc.).
   - Distribute classwork problems and homework problems into the matching individual lecture unit (2 to 4 problems per unit).
3. If available spreadsheet tabs are provided, match each course to the best corresponding tab name.

Output schema: Return a JSON array of course objects:
[
  {{
    "course_name": "Full Subject Title",
    "course_code": "SHORT_CODE",
    "matched_tab": "Matching Tab Name",
    "units": [
      {{
        "unit_id": "SHORT_CODE-1",
        "week": 1,
        "topic": "Concise Lecture Topic Name",
        "content": ["Subtopic 1", "Subtopic 2"],
        "classwork": ["Problem 1", "Problem 2"],
        "homework": ["Problem 1", "Problem 2"]
      }}
    ]
  }}
]

CRITICAL RULES:
1. Every individual lecture/class must be its own separate unit.
2. Return strictly valid JSON array `[ ... ]`. No conversational wrapper.
3. Even if only 1 subject is present, return it inside the array `[ {{ ... }} ]`.
4. Number units sequentially for each course starting from 1 (e.g. DAA-1, DAA-2, ...).
"""

    models_to_try = [groq_model, "openai/gpt-oss-120b", "openai/gpt-oss-20b", "qwen/qwen3.8-27b"]
    seen_models = set()
    response = None

    for m in models_to_try:
        if not m or m in seen_models:
            continue
        seen_models.add(m)
        tok_limit = 1000 if "qwen" in m else 2500
        try:
            response = client.chat.completions.create(
                model=m,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": f"Extract all courses and units from this document:\n\n{cleaned_text}"},
                ],
                temperature=0.1,
                max_tokens=tok_limit,
            )
            break
        except Exception as ge:
            print(f"[WARN] Multi-course model {m} failed: {ge}, trying fallback...")
            continue

    if not response:
        raise RuntimeError("All Groq models encountered rate limits or errors.")

    content = response.choices[0].message.content or ""
    clean_json_str = content.strip()
    if "```json" in clean_json_str:
        clean_json_str = clean_json_str.split("```json")[1].split("```")[0].strip()
    elif "```" in clean_json_str:
        clean_json_str = clean_json_str.split("```")[1].split("```")[0].strip()

    try:
        courses = json.loads(clean_json_str)
    except Exception as e:
        match = re.search(r"(\[.*\]|\{.*\})", clean_json_str, re.DOTALL)
        if match:
            courses = json.loads(match.group(1))
        else:
            raise ValueError(f"Could not parse multi-course output: {e}\n{content[:200]}")

    if isinstance(courses, dict):
        if "courses" in courses:
            courses = courses["courses"]
        else:
            courses = [courses]

    normalized_courses = []
    for c in courses:
        c_name = c.get("course_name") or "Course"
        c_code = c.get("course_code") or "CRS"
        m_tab = c.get("matched_tab") or c_code
        raw_units = c.get("units") or []

        norm_units = []
        for idx, u in enumerate(raw_units, start=1):
            uid = u.get("unit_id") or f"{c_code}-{idx}"
            top = u.get("topic") or f"Module {idx}"
            wk = u.get("week") or ((idx - 1) // 3 + 1)
            cnt = u.get("content") or []
            if isinstance(cnt, str):
                cnt = [x.strip() for x in cnt.split("\n") if x.strip()]
            cw = u.get("classwork") or []
            if isinstance(cw, str):
                cw = [x.strip() for x in cw.split("\n") if x.strip()]
            hw = u.get("homework") or []
            if isinstance(hw, str):
                hw = [x.strip() for x in hw.split("\n") if x.strip()]

            norm_units.append({
                "unit_id": str(uid).strip(),
                "week": int(wk) if str(wk).isdigit() else 1,
                "topic": str(top).strip(),
                "content": cnt,
                "classwork": cw,
                "homework": hw,
            })

        normalized_courses.append({
            "course_name": str(c_name).strip(),
            "course_code": str(c_code).strip().upper(),
            "matched_tab": str(m_tab).strip(),
            "units": norm_units,
        })

    return normalized_courses

